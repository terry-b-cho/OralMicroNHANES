#!/usr/bin/env python3
"""Assemble the styled supplementary workbook from staged CSVs.

Called by 09_supplementary_excel.R with:  <manifest.json> <output.xlsx>
Manifest: [{"sheet": name, "csv": path, "long_text": bool}, ...] in order.
Styling: bold white-on-slate frozen header, autofilter, capped auto column
widths, Arial 9 body font (Excel display; unrelated to figure font rules).
"""
import csv, json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

manifest_path, out_path = sys.argv[1], sys.argv[2]
with open(manifest_path) as fh:
    manifest = json.load(fh)

HDR_FILL = PatternFill("solid", fgColor="44546A")
HDR_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=9)
WRAP = Alignment(wrap_text=True, vertical="top")

def maybe_number(v):
    if v is None or v == "":
        return None
    if v in ("TRUE", "FALSE"):
        return v == "TRUE"
    try:
        f = float(v)
        return int(f) if f.is_integer() and "e" not in v.lower() and "." not in v else f
    except ValueError:
        return v

wb = Workbook()
wb.remove(wb.active)
for entry in manifest:
    ws = wb.create_sheet(title=entry["sheet"][:31])
    with open(entry["csv"], newline="") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        continue
    widths = [0] * len(rows[0])
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row, start=1):
            cell = ws.cell(row=r_i, column=c_i,
                           value=val if r_i == 1 else maybe_number(val))
            if r_i == 1:
                cell.font, cell.fill = HDR_FONT, HDR_FILL
                cell.alignment = Alignment(vertical="center")
            else:
                cell.font = BODY_FONT
                if entry.get("long_text"):
                    cell.alignment = WRAP
            if c_i <= len(widths):
                widths[c_i - 1] = max(widths[c_i - 1], min(len(str(val)), 80))
    for c_i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c_i)].width = min(max(w + 2, 9),
            60 if entry.get("long_text") else 32)
    ws.freeze_panes = "A2"
    if not entry.get("long_text") and len(rows) > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"
wb.save(out_path)
print(f"Workbook written: {out_path} ({len(manifest)} sheets)")
